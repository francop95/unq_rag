"""
Extrae el glosario de nomenclatura del inventario
=================================================

La hoja `Nomenclatura` del Excel mapea cada componente a su código acordado
(QD01, SSR01, VFD01…), su nombre técnico y los alias con los que aparece en los
planos y en la tesis. Es la tabla que permite que una respuesta diga "el
variador (VFD01)" en vez de "el variador" a secas.

Se vuelca a JSON en vez de leer el Excel desde la API por dos razones: la
imagen de la API no lleva openpyxl, y un JSON versionado se puede revisar en un
diff cuando el inventario cambie, que es cuando conviene mirarlo.

    python scripts/generar_nomenclatura.py
"""

import json
import os
import sys

import openpyxl

RAIZ = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
EXCEL = os.path.join(RAIZ, "data", "raw_data", "Inventario_secadero_revision_Pablo_v2.xlsx")
SALIDA = os.path.abspath(os.path.join(RAIZ, "..", "API", "configs", "nomenclatura.json"))


def main() -> int:
    if not os.path.exists(EXCEL):
        print(f"  no está el inventario: {EXCEL}")
        return 1

    wb = openpyxl.load_workbook(EXCEL, data_only=True)
    if "Nomenclatura" not in wb.sheetnames:
        print("  el inventario no tiene hoja 'Nomenclatura'")
        return 1

    # Ubicación física por código, desde el inventario: la hoja de nomenclatura
    # tiene la referencia al plano pero no dónde está montado el componente.
    ubicaciones = {}
    if "Inventario" in wb.sheetnames:
        inv = wb["Inventario"]
        fs = [[("" if c is None else str(c).strip()) for c in r]
              for r in inv.iter_rows(values_only=True)]
        fs = [f for f in fs if any(f)]
        for f in fs[3:]:
            cod = f[10].strip() if len(f) > 10 and f[10] else ""
            if cod:
                ubicaciones[cod] = f[7].strip() if len(f) > 7 and f[7] else ""

    ws = wb["Nomenclatura"]
    filas = [[("" if c is None else str(c).strip()) for c in r]
             for r in ws.iter_rows(values_only=True)]
    filas = [f for f in filas if any(f)]

    # Dos filas de título y una de encabezados; los datos empiezan en la cuarta.
    componentes = []
    for f in filas[3:]:
        codigo = (f[0] if len(f) > 0 else "").strip()
        nombre = (f[2] if len(f) > 2 else "").strip()
        alias = (f[3] if len(f) > 3 else "").strip()
        if not codigo or not nombre:
            continue
        componentes.append({
            "codigo": codigo, "nombre": nombre, "alias": alias,
            # Dónde encontrarlo: la referencia en el plano rotulado ("P1: ramal
            # calefacción") y la ubicación física ("Tablero"). Sin esto la
            # respuesta da el código pero no dice dónde ir a buscarlo, que es la
            # mitad del trabajo de mantenimiento.
            "en_plano": (f[4] if len(f) > 4 else "").strip(),
            "ubicacion": ubicaciones.get(codigo, ""),
        })

    os.makedirs(os.path.dirname(SALIDA), exist_ok=True)
    with open(SALIDA, "w", encoding="utf-8") as fh:
        json.dump({"componentes": componentes}, fh, ensure_ascii=False, indent=2)

    print(f"  {len(componentes)} componentes → {SALIDA}")
    for c in componentes[:4]:
        print(f"    {c['codigo']:<8} {c['nombre'][:46]}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
